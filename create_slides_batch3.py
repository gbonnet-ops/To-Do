from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

DARK_BLUE = RGBColor(0x15, 0x23, 0x4A)
LIGHT_BLUE = RGBColor(0x65, 0xC4, 0xDA)
LIGHT_GREEN = RGBColor(0xB6, 0xCF, 0x2B)
GREEN = RGBColor(0x95, 0xCD, 0x9C)
GREY = RGBColor(0xE5, 0xE5, 0xE5)
DARK_BLUE2 = RGBColor(0x2A, 0x38, 0x5B)
BG_COLOR = RGBColor(0xFA, 0xFA, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)

def add_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG_COLOR

def add_header(slide, section, title, page_num):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), prs.slide_width, Cm(1.2))
    shp.fill.solid(); shp.fill.fore_color.rgb = DARK_BLUE; shp.line.fill.background()
    tf = shp.text_frame; p = tf.paragraphs[0]
    p.text = section; p.font.color.rgb = WHITE; p.font.size = Pt(9); p.font.bold = True
    run = p.add_run(); run.text = f"    |    {page_num}"; run.font.color.rgb = LIGHT_BLUE; run.font.size = Pt(9)
    txBox = slide.shapes.add_textbox(Cm(1.5), Cm(1.5), Cm(30), Cm(1.5))
    tf = txBox.text_frame; p = tf.paragraphs[0]; p.text = title
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = DARK_BLUE

def add_footer(slide):
    txBox = slide.shapes.add_textbox(Cm(1), Cm(18), Cm(12), Cm(0.8))
    p = txBox.text_frame.paragraphs[0]; p.text = "Strictly Private and Confidential"
    p.font.size = Pt(7); p.font.color.rgb = RGBColor(0x99, 0x99, 0x99); p.font.italic = True

def add_kpi_box(slide, left, top, w, h, val, label, color=LIGHT_BLUE):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = color; shp.line.width = Pt(1.5)
    tf = shp.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = val; p.font.size = Pt(20); p.font.bold = True
    p.font.color.rgb = DARK_BLUE; p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph(); p2.text = label; p2.font.size = Pt(8)
    p2.font.color.rgb = DARK_BLUE2; p2.alignment = PP_ALIGN.CENTER

# ============================================
# SLIDE 23: Unique asset at crossroads
# ============================================
s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
add_header(s, "Product & Services", "GMSL is a unique asset at the crossroads of energy operations and software", 23)

# Revenue breakdown boxes
segments = [
    ("Operations", "58.3%", "£10.2m", "Shipper Ops + TSO Ops\n24/7 outsourced dispatching", LIGHT_BLUE),
    ("Software", "33.8%", "£5.9m", "Chorus, EuroRunner, CodeRunner\nPowerTrak, ENOM", LIGHT_GREEN),
    ("CVA", "7.9%", "£1.4m", "Sole Claims Validation Agent\nfor UK beach terminals", GREEN),
]
for i, (name, pct, rev, desc, color) in enumerate(segments):
    left = Cm(1.5 + i * 10.5)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Cm(3.5), Cm(9.8), Cm(7))
    box.fill.solid(); box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = color; box.line.width = Pt(2)
    tf = box.text_frame; tf.word_wrap = True
    # Color bar at top
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Cm(3.5), Cm(9.8), Cm(0.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
    tf_bar = bar.text_frame; p = tf_bar.paragraphs[0]
    p.text = name; p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
    
    # Percentage
    p = tf.paragraphs[0]; p.text = pct; p.font.size = Pt(32); p.font.bold = True
    p.font.color.rgb = color; p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(12)
    # Revenue
    p2 = tf.add_paragraph(); p2.text = rev; p2.font.size = Pt(14); p2.font.bold = True
    p2.font.color.rgb = DARK_BLUE; p2.alignment = PP_ALIGN.CENTER; p2.space_after = Pt(8)
    # Description
    p3 = tf.add_paragraph(); p3.text = desc; p3.font.size = Pt(9)
    p3.font.color.rgb = DARK_BLUE2; p3.alignment = PP_ALIGN.CENTER

# Bottom section - key strengths
bottom = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1.5), Cm(11.5), Cm(31), Cm(6))
bottom.fill.solid(); bottom.fill.fore_color.rgb = WHITE
bottom.line.color.rgb = GREY; bottom.line.width = Pt(0.75)
tf = bottom.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.text = "FY26F Total Revenue: £17,557k"
p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = DARK_BLUE; p.space_after = Pt(6)
items = [
    "Combined software + operations offering creates unique value proposition in European energy scheduling",
    "Mission-critical 24/7 services drive deep client integration and high switching costs",
    "Growing recurring revenue base with annual inflation indexation and evergreen contracts",
    "Lean cost structure with staff costs at 83.5% of total OpEx",
]
for item in items:
    p = tf.add_paragraph(); p.text = f"• {item}"; p.font.size = Pt(9); p.font.color.rgb = DARK_BLUE2; p.space_after = Pt(3)

note = s.shapes.add_textbox(Cm(1.5), Cm(17.2), Cm(20), Cm(0.8))
p = note.text_frame.paragraphs[0]
p.text = "[Sources: Slide 6 – GMSL Asset Fiche / Screenshot Grégoire]"
p.font.size = Pt(7); p.font.italic = True; p.font.color.rgb = RGBColor(0x99,0x99,0x99)
add_footer(s)

# ============================================
# SLIDE 24: Europe's leading outsourced dispatching
# ============================================
s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
add_header(s, "Product & Services", "Europe's leading outsourced energy dispatching service", 24)

# KPI row
kpis = [
    ("24/7/365", "Dispatching Operations\nfrom Cambridge", LIGHT_BLUE),
    ("60+", "Operations Staff\n(Shipper + TSO Ops)", LIGHT_GREEN),
    ("90+", "Client Companies\nacross Europe", LIGHT_BLUE),
    ("200+", "TSO/SSO\nCounterparties", LIGHT_GREEN),
    ("Multiple", "Gas Networks &\nPower Grids Covered", LIGHT_BLUE),
]
for i, (val, label, color) in enumerate(kpis):
    add_kpi_box(s, Cm(1.5 + i*6.3), Cm(3.3), Cm(5.8), Cm(2.8), val, label, color)

# Left - Shipper Ops
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1.5), Cm(7), Cm(15), Cm(10.5))
box.fill.solid(); box.fill.fore_color.rgb = WHITE
box.line.color.rgb = LIGHT_BLUE; box.line.width = Pt(1)
tf = box.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.text = "Shipper Operations (£7.5m FY26F)"
p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = DARK_BLUE; p.space_after = Pt(6)
items = [
    "Day-ahead position import from ETRM systems",
    "Automated nomination to TSOs via Edig@s / AS2 / AS4",
    "Real-time matching and balancing display",
    "Within-day re-nominations",
    "Post-day allocations and reporting",
    "Portfolio management across multiple networks",
    "Clients: Shell, BP, TotalEnergies, Vitol, Trafigura, Centrica, ExxonMobil, Glencore, and 70+ others",
]
for item in items:
    p = tf.add_paragraph(); p.text = f"• {item}"; p.font.size = Pt(9); p.font.color.rgb = DARK_BLUE2; p.space_after = Pt(3)

# Right - TSO Ops
box2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(17.5), Cm(7), Cm(15), Cm(10.5))
box2.fill.solid(); box2.fill.fore_color.rgb = WHITE
box2.line.color.rgb = LIGHT_GREEN; box2.line.width = Pt(1)
tf = box2.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.text = "TSO Operations (£2.7m FY26F)"
p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = DARK_BLUE; p.space_after = Pt(6)
items2 = [
    "Physically separate team for compliance",
    "15 dedicated TSO operators (1 Mgr, 2 TLs, 6 Senior, 6 Junior)",
    "Network monitoring and balancing operations",
    "Capacity booking and allocation management",
    "DR facility in Cambridge with 24/7 IT cover",
    "Hot-standby site for business continuity",
    "Dedicated coverage for networks, hardware and communications",
]
for item in items2:
    p = tf.add_paragraph(); p.text = f"• {item}"; p.font.size = Pt(9); p.font.color.rgb = DARK_BLUE2; p.space_after = Pt(3)

note = s.shapes.add_textbox(Cm(1.5), Cm(17.2), Cm(20), Cm(0.8))
p = note.text_frame.paragraphs[0]
p.text = "[Sources: GMSL Asset Fiche p.6, p.9 / gmsl.co.uk/operations]"
p.font.size = Pt(7); p.font.italic = True; p.font.color.rgb = RGBColor(0x99,0x99,0x99)
add_footer(s)

# ============================================
# SLIDE 25: Software suite - full scheduling workflow
# ============================================
s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
add_header(s, "Product & Services", "A purpose-built software suite covering the full scheduling workflow", 25)

products = [
    ("CodeRunner", "Gas nomination & scheduling for UK networks", "UK gas (NBP, Fluxys)", 
     ["Automated gas nominations", "TSO connectivity", "Real-time balancing", "Position management"], LIGHT_BLUE),
    ("EuroRunner", "Multi-network gas portfolio management", "European gas hubs (Benelux, DE)", 
     ["Multi-network portfolios", "Edig@s v6.1 upgrade planned", "API modernisation", "PRISMA interface"], LIGHT_GREEN),
    ("ENOM / ENOM+", "Energy nomination workflow", "UK & European gas", 
     ["Nomination workflows", "Matching & allocation", "Client migration to Chorus", "Legacy platform"], LIGHT_BLUE),
    ("PowerTrak", "Power scheduling with ETRM integration", "European power markets", 
     ["Web-based interface", "ETRM integration", "Central/Eastern EU expansion", "Power scheduling"], LIGHT_GREEN),
    ("Chorus", "Modern cloud-native SaaS platform", "Italy (PSV), Spain, Central EU", 
     ["Cloud-native (AWS)", "Modern SaaS UI", "Next-gen platform", "Replacing legacy products"], LIGHT_BLUE),
]

for i, (name, overview, market, features, color) in enumerate(products):
    left = Cm(1 + i * 6.5)
    # Product card
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Cm(3.3), Cm(6.1), Cm(14.2))
    box.fill.solid(); box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = color; box.line.width = Pt(1.5)
    
    # Name header bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Cm(3.3), Cm(6.1), Cm(1.2))
    bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
    tf = bar.text_frame; p = tf.paragraphs[0]; p.text = name
    p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
    
    tf = box.text_frame; tf.word_wrap = True
    # Overview
    p = tf.paragraphs[0]; p.space_before = Pt(20)
    p.text = "Overview"; p.font.size = Pt(8); p.font.bold = True; p.font.color.rgb = DARK_BLUE
    p2 = tf.add_paragraph(); p2.text = overview; p2.font.size = Pt(8); p2.font.color.rgb = DARK_BLUE2; p2.space_after = Pt(6)
    # Market
    p3 = tf.add_paragraph(); p3.text = "Market Coverage"; p3.font.size = Pt(8); p3.font.bold = True; p3.font.color.rgb = DARK_BLUE
    p4 = tf.add_paragraph(); p4.text = market; p4.font.size = Pt(8); p4.font.color.rgb = DARK_BLUE2; p4.space_after = Pt(6)
    # Features
    p5 = tf.add_paragraph(); p5.text = "Key Features"; p5.font.size = Pt(8); p5.font.bold = True; p5.font.color.rgb = DARK_BLUE
    for feat in features:
        p6 = tf.add_paragraph(); p6.text = f"• {feat}"; p6.font.size = Pt(8); p6.font.color.rgb = DARK_BLUE2; p6.space_after = Pt(2)
    # Delivery
    p7 = tf.add_paragraph(); p7.space_before = Pt(6); p7.text = "Delivery Model"; p7.font.size = Pt(8); p7.font.bold = True; p7.font.color.rgb = DARK_BLUE
    p8 = tf.add_paragraph()
    p8.text = "SaaS / Hosted" if name == "Chorus" else "On-premise / Hosted"
    p8.font.size = Pt(8); p8.font.color.rgb = DARK_BLUE2
    # Launch date
    p9 = tf.add_paragraph(); p9.space_before = Pt(6); p9.text = "Launch Date"; p9.font.size = Pt(8); p9.font.bold = True; p9.font.color.rgb = DARK_BLUE
    p10 = tf.add_paragraph(); p10.text = "[See Fluxys Timeline]"; p10.font.size = Pt(7); p10.font.italic = True; p10.font.color.rgb = RGBColor(0x99,0x99,0x99)

note = s.shapes.add_textbox(Cm(1.5), Cm(17.3), Cm(25), Cm(0.7))
p = note.text_frame.paragraphs[0]
p.text = "[Management to confirm launch date of Chorus | Sources: Asset Fiche p.10, Fluxys Timeline in 4. Additional Materials]"
p.font.size = Pt(7); p.font.italic = True; p.font.color.rgb = RGBColor(0x99,0x99,0x99)
add_footer(s)

prs.save("/tmp/darwin_slides_batch3.pptx")
print("Batch 3 saved: Slides 23, 24, 25")
