from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Darwin theme colors
DARK_BLUE = RGBColor(0x15, 0x23, 0x4A)
LIGHT_BLUE = RGBColor(0x65, 0xC4, 0xDA)
LIGHT_GREEN = RGBColor(0xB6, 0xCF, 0x2B)
GREEN = RGBColor(0x95, 0xCD, 0x9C)
GREY = RGBColor(0xE5, 0xE5, 0xE5)
DARK_BLUE2 = RGBColor(0x2A, 0x38, 0x5B)
BG_COLOR = RGBColor(0xFA, 0xFA, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)

def add_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, section, title, page_num):
    # Section header bar
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), prs.slide_width, Cm(1.2))
    shp.fill.solid()
    shp.fill.fore_color.rgb = DARK_BLUE
    shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = section
    p.font.color.rgb = WHITE
    p.font.size = Pt(9)
    p.font.bold = True
    # Page number
    run = p.add_run()
    run.text = f"    |    {page_num}"
    run.font.color.rgb = LIGHT_BLUE
    run.font.size = Pt(9)
    
    # Title
    txBox = slide.shapes.add_textbox(Cm(1.5), Cm(1.5), Cm(30), Cm(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

def add_footer(slide):
    txBox = slide.shapes.add_textbox(Cm(1), Cm(18), Cm(12), Cm(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Strictly Private and Confidential"
    p.font.size = Pt(7)
    p.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    p.font.italic = True

def add_kpi_box(slide, left, top, width, height, value, label, color=LIGHT_BLUE):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = color
    shp.line.width = Pt(1.5)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    # Value
    p = tf.paragraphs[0]
    p.text = value
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    # Label
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(9)
    p2.font.color.rgb = DARK_BLUE2
    p2.alignment = PP_ALIGN.CENTER

def add_bullet_list(tf, items, font_size=Pt(10), color=DARK_BLUE):
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = font_size
        p.font.color.rgb = color
        p.space_after = Pt(4)

# ============================================
# SLIDE 5: Opportunity overview & post-transaction structure
# ============================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide5)
add_header(slide5, "Executive summary", "Opportunity overview and illustrative post-transaction structure", 5)

# Left column - Opportunity overview
box_left = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1.5), Cm(3.5), Cm(15), Cm(14))
box_left.fill.solid()
box_left.fill.fore_color.rgb = WHITE
box_left.line.color.rgb = GREY
box_left.line.width = Pt(0.75)
tf = box_left.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Opportunity Overview"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = DARK_BLUE
p.space_after = Pt(12)

items = [
    "GMSL is a unique and resilient asset combining 24/7 outsourced dispatching operations, mission-critical SaaS platforms, and UK CVA services",
    "Serving 130+ clients across multiple gas networks and power grids in Europe",
    "Strategic combination with Energy One to create a leading European energy market software and operations platform",
    "Highly recurring revenue model (software + managed services) with strong earnings visibility",
    "FY26F total revenue: £17,557k (Operations 58.3%, Software 33.8%)",
    "Structural growth driven by energy transition, rising power demand, and digitalisation of power markets",
]
add_bullet_list(tf, items)

# Right column - Post-transaction structure placeholder
box_right = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(17.5), Cm(3.5), Cm(15), Cm(14))
box_right.fill.solid()
box_right.fill.fore_color.rgb = WHITE
box_right.line.color.rgb = GREY
box_right.line.width = Pt(0.75)
tf = box_right.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Illustrative Post-Transaction Structure"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = DARK_BLUE
p.space_after = Pt(12)

items2 = [
    "Combined entity: GMSL + Energy One",
    "Leading European energy market software & operations platform",
    "Complementary geographic coverage and product suites",
    "Enhanced cross-selling opportunities across client bases",
    "Accelerated commercial growth and expanded European footprint",
]
add_bullet_list(tf, items2)

# Note box
note = slide5.shapes.add_textbox(Cm(17.5), Cm(15), Cm(15), Cm(2))
tf = note.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "[Org chart / structure diagram to be inserted]"
p.font.size = Pt(9)
p.font.italic = True
p.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

add_footer(slide5)

# ============================================
# SLIDE 6: Key Investment Highlights
# ============================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide6)
add_header(slide6, "Executive summary", "Key Investment Highlights", 6)

highlights = [
    ("1", "Mission-Critical\nPosition", "24/7 outsourced dispatching at the heart of European energy flows. TSO connectivity to 200+ counterparties creates deep integration and high switching costs", LIGHT_BLUE),
    ("2", "Highly Recurring\nRevenue", "Software & managed services with annual contracts, automatic renewal, and inflation indexation. Strong earnings visibility with consistent cash conversion", LIGHT_GREEN),
    ("3", "Structural Market\nGrowth", "Rising power demand, infrastructure expansion, energy transition (H2, CO2), increasing market participants, and digitalisation drive expanding TAM", LIGHT_BLUE),
    ("4", "Unique Competitive\nMoat", "30 years of domain expertise, certified TSO connections, 24/7 operational capability, and regulatory mandates create formidable barriers to entry", LIGHT_GREEN),
    ("5", "Attractive\nFinancial Profile", "FY26F revenue £17.6m with lean cost structure (83.5% staff costs). Target: £3.2m PAT + inflation. Growing subscription-based revenues", LIGHT_BLUE),
    ("6", "Strategic\nCombination", "Combination with Energy One creates leading European energy market platform, accelerating growth through geographic expansion and cross-selling", LIGHT_GREEN),
]

for i, (num, title, desc, color) in enumerate(highlights):
    col = i % 3
    row = i // 3
    left = Cm(1.5 + col * 10.5)
    top = Cm(3.5 + row * 7.5)
    
    # Number circle
    circle = slide6.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Cm(1.2), Cm(1.2))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    tf = circle.text_frame
    tf.paragraphs[0].text = num
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Title
    txBox = slide6.shapes.add_textbox(left + Cm(1.5), top - Cm(0.2), Cm(8.5), Cm(1.4))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    
    # Description
    txBox2 = slide6.shapes.add_textbox(left, top + Cm(1.5), Cm(9.5), Cm(5))
    tf = txBox2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = desc
    p.font.size = Pt(9)
    p.font.color.rgb = DARK_BLUE2
    p.space_after = Pt(4)

add_footer(slide6)

# ============================================
# SLIDE 15: Market size
# ============================================
slide15 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide15)
add_header(slide15, "Market", "Market size", 15)

# KPI boxes at top
kpis = [
    ("£17.6m", "GMSL FY26F Revenue"),
    ("139", "Client Relationships"),
    ("200+", "TSO Counterparties"),
    ("90+", "Operations Clients"),
    ("100+", "Software Subscriptions"),
]
for i, (val, label) in enumerate(kpis):
    add_kpi_box(slide15, Cm(1.5 + i * 6.3), Cm(3.3), Cm(5.8), Cm(2.5), val, label, 
                LIGHT_BLUE if i % 2 == 0 else LIGHT_GREEN)

# Left section - Growth drivers
box = slide15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1.5), Cm(6.5), Cm(15), Cm(11))
box.fill.solid()
box.fill.fore_color.rgb = WHITE
box.line.color.rgb = GREY
box.line.width = Pt(0.75)
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Structural Growth Drivers"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = DARK_BLUE
p.space_after = Pt(8)

drivers = [
    "Growing power demand across Europe",
    "Infrastructure build-out in gas and power networks",
    "Increasing variability in power generation mix (renewables)",
    "Rising number of market participants",
    "Energy transition creating new commodities (H₂, CO₂) and new players",
    "Digitalisation and automation of power markets",
    "High barriers to entry: specialised knowledge, 24/7 requirements, certified TSO connectivity",
    "Typical software contract values of £20k-£150k/year",
]
add_bullet_list(tf, drivers, Pt(9))

# Right section - Geographic expansion
box2 = slide15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(17.5), Cm(6.5), Cm(15), Cm(11))
box2.fill.solid()
box2.fill.fore_color.rgb = WHITE
box2.line.color.rgb = GREY
box2.line.width = Pt(0.75)
tf = box2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Market Coverage & Expansion"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = DARK_BLUE
p.space_after = Pt(8)

coverage = [
    "Strong presence: UK, Benelux, Germany, key European gas hubs",
    "Chorus expansion: Italy (PSV confirmed, transit next), Spain, Central Europe",
    "PowerTrak: expanding towards Central/Eastern Europe",
    "EuroRunner: API modernisation and Edig@s v6.1 + PRISMA interface planned",
    "90+ operations clients and 100+ software subscriptions = fraction of total European gas & power participants",
    "Significant whitespace in addressable market remains",
]
add_bullet_list(tf, coverage, Pt(9))

# Chart placeholder
note = slide15.shapes.add_textbox(Cm(17.5), Cm(15.5), Cm(15), Cm(1.5))
tf = note.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "[TAM chart / market sizing graph to be inserted from Infopack]"
p.font.size = Pt(9)
p.font.italic = True
p.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

add_footer(slide15)

# Save
prs.save("/tmp/darwin_slides_batch1.pptx")
print("Batch 1 saved: Slides 5, 6, 15")
